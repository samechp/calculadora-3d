import webview
import json
import os
import sys
import time
import shutil
import threading
import tempfile
import subprocess
import webbrowser
import urllib.request
from fpdf import FPDF
from openpyxl import Workbook
import datetime

def get_base_path():
    if getattr(sys, 'frozen', False):
        return os.path.dirname(sys.executable)
    return os.path.dirname(os.path.abspath(__file__))

PROFILES_FILE = os.path.join(get_base_path(), 'perfiles.json')

# fpdf 1.7 solo sabe escribir latin-1. Si el usuario escribe el nombre de su
# empresa con comillas tipográficas, un guion largo o un emoji, la exportación
# entera reventaría. Se cambian los casos comunes y se descarta el resto.
_CAMBIOS_PDF = {
    '‘': "'", '’': "'", '“': '"', '”': '"',
    '–': '-', '—': '-', '…': '...', ' ': ' ',
    '•': '-', '€': 'EUR',
}


def _texto_pdf(texto):
    if texto is None:
        return ''
    texto = str(texto)
    for malo, bueno in _CAMBIOS_PDF.items():
        texto = texto.replace(malo, bueno)
    return texto.encode('latin-1', 'ignore').decode('latin-1')


# ===== COPIAS DE SEGURIDAD =====
# Todo (perfiles, piezas, proyectos, filamentos) vive en un solo archivo que se
# sobrescribe en cada guardado. Antes de pisarlo se guarda una copia, y se
# conservan las últimas COPIAS_MAX por si hay que volver atrás.
COPIAS_DIR = os.path.join(get_base_path(), 'copias')
COPIAS_MAX = 10


def _copias_ordenadas():
    if not os.path.isdir(COPIAS_DIR):
        return []
    archivos = [f for f in os.listdir(COPIAS_DIR) if f.startswith('perfiles_') and f.endswith('.json')]
    archivos.sort(reverse=True)   # el nombre lleva la fecha, así que ordena solo
    return archivos


def hacer_copia():
    """Guarda el estado anterior antes de sobrescribirlo. Si algo falla, no se
    interrumpe el guardado: perder una copia es mucho menos grave que no guardar."""
    try:
        if not os.path.exists(PROFILES_FILE) or os.path.getsize(PROFILES_FILE) < 10:
            return
        if not os.path.isdir(COPIAS_DIR):
            os.makedirs(COPIAS_DIR)

        nombre = 'perfiles_{}.json'.format(datetime.datetime.now().strftime('%Y-%m-%d_%H-%M-%S'))
        destino = os.path.join(COPIAS_DIR, nombre)
        if os.path.exists(destino):
            return   # ya hay una copia de este mismo segundo
        shutil.copy2(PROFILES_FILE, destino)

        for viejo in _copias_ordenadas()[COPIAS_MAX:]:
            try:
                os.remove(os.path.join(COPIAS_DIR, viejo))
            except Exception:
                pass
    except Exception as e:
        print("No se pudo hacer la copia de seguridad:", e)
# ===== FIN COPIAS DE SEGURIDAD =====

# ===== TAMAÑO Y POSICIÓN DE LA VENTANA (se recuerdan entre sesiones) =====
WINDOW_FILE = os.path.join(get_base_path(), 'ventana.json')

# Primera vez que se abre: ventana grande y maximizada
DEFAULT_WINDOW = {'width': 1400, 'height': 950, 'x': None, 'y': None, 'maximized': True}

MIN_W, MIN_H = 900, 600
MAX_W, MAX_H = 20000, 20000


def _virtual_screen():
    """Rectángulo que abarca todos los monitores (en píxeles físicos)."""
    try:
        import ctypes
        gsm = ctypes.windll.user32.GetSystemMetrics
        x, y = gsm(76), gsm(77)          # SM_XVIRTUALSCREEN / SM_YVIRTUALSCREEN
        w, h = gsm(78), gsm(79)          # SM_CXVIRTUALSCREEN / SM_CYVIRTUALSCREEN
        if w > 0 and h > 0:
            return x, y, w, h
    except Exception:
        pass
    return None


def load_window_state():
    st = dict(DEFAULT_WINDOW)
    try:
        if os.path.exists(WINDOW_FILE):
            with open(WINDOW_FILE, 'r', encoding='utf-8') as f:
                data = json.load(f)
            if isinstance(data, dict):
                for k in st:
                    if k in data and data[k] is not None:
                        st[k] = data[k]
    except Exception as e:
        print("Error cargando estado de ventana:", e)

    # Saneamiento: medidas absurdas vuelven al valor por defecto
    try:
        st['width'] = min(MAX_W, max(MIN_W, int(st['width'])))
        st['height'] = min(MAX_H, max(MIN_H, int(st['height'])))
    except Exception:
        st['width'], st['height'] = DEFAULT_WINDOW['width'], DEFAULT_WINDOW['height']

    st['maximized'] = bool(st.get('maximized'))

    # Si la posición guardada quedó fuera de los monitores actuales, se centra
    try:
        st['x'] = None if st['x'] is None else int(st['x'])
        st['y'] = None if st['y'] is None else int(st['y'])
    except Exception:
        st['x'] = st['y'] = None

    if st['x'] is not None and st['y'] is not None:
        vs = _virtual_screen()
        if vs:
            vx, vy, vw, vh = vs
            visible = (st['x'] + st['width'] > vx + 80 and st['x'] < vx + vw - 80
                       and st['y'] + 40 > vy and st['y'] < vy + vh - 40)
            if not visible:
                st['x'] = st['y'] = None

    return st


def save_window_state(st):
    try:
        with open(WINDOW_FILE, 'w', encoding='utf-8') as f:
            json.dump(st, f)
    except Exception as e:
        print("Error guardando estado de ventana:", e)
# ===== FIN TAMAÑO Y POSICIÓN DE LA VENTANA =====


# ===== ACTUALIZACIÓN AUTOMÁTICA =====
# Hay dos cosas que se pueden actualizar por separado:
#   - La INTERFAZ (html/css/js): son unos pocos KB. Se descarga sola en segundo
#     plano y queda lista para el siguiente arranque.
#   - El EJECUTABLE (esta parte en Python): son ~34 MB. Solo se avisa; el usuario
#     decide con un botón y la app se reinicia sola.
APP_VERSION = '1.0.3'   # versión del .exe
WEB_VERSION = '1.0.7'   # versión de la interfaz que viene dentro del .exe

REPO = 'samechp/calculadora-3d'
VERSION_URL = 'https://raw.githubusercontent.com/{}/main/version.json'.format(REPO)
WEB_BASE_URL = 'https://raw.githubusercontent.com/{}/main/web_app/'.format(REPO)
EXE_URL = 'https://github.com/{}/releases/latest/download/Calculadora3D.exe'.format(REPO)
RELEASES_URL = 'https://github.com/{}/releases/latest'.format(REPO)

RELEASES_API = 'https://api.github.com/repos/{}/releases'.format(REPO)

def url_exe_de_version(version):
    return 'https://github.com/{}/releases/download/v{}/Calculadora3D.exe'.format(REPO, version)

WEB_UPDATE_DIR = os.path.join(get_base_path(), 'web_update')
WEB_UPDATE_VERSION_FILE = os.path.join(WEB_UPDATE_DIR, 'version.txt')
PREFS_FILE = os.path.join(get_base_path(), 'actualizaciones.json')

# Estado compartido con la interfaz mientras se baja el ejecutable
_estado_update = {'estado': 'inactivo', 'porcentaje': 0, 'mensaje': ''}
_info_remota = {}


def cargar_prefs():
    """omitidas: versiones que el usuario no quiere que le vuelvan a ofrecer.
    fijada: si el usuario se quedó a propósito en una versión vieja, no se
    actualiza nada (ni interfaz ni ejecutable) hasta que él lo reactive."""
    prefs = {'omitidas': [], 'fijada': None}
    try:
        if os.path.exists(PREFS_FILE):
            with open(PREFS_FILE, 'r', encoding='utf-8') as f:
                data = json.load(f)
            if isinstance(data, dict):
                prefs['omitidas'] = [str(v) for v in data.get('omitidas', [])]
                prefs['fijada'] = data.get('fijada') or None
    except Exception as e:
        print("Error leyendo preferencias de actualización:", e)
    return prefs


def guardar_prefs(prefs):
    try:
        with open(PREFS_FILE, 'w', encoding='utf-8') as f:
            json.dump(prefs, f, ensure_ascii=False)
    except Exception as e:
        print("Error guardando preferencias de actualización:", e)


def _numero_version(v):
    """'2.10.1' -> (2, 10, 1). Sirve para comparar versiones de verdad,
    no como texto (donde '2.9' saldría mayor que '2.10')."""
    try:
        return tuple(int(p) for p in str(v).strip().split('.'))
    except Exception:
        return (0,)


def es_mas_nueva(remota, local):
    return _numero_version(remota) > _numero_version(local)


def version_web_actual():
    """Versión de interfaz que se va a usar: la descargada si es más nueva
    que la que trae el .exe, si no la incluida."""
    try:
        if os.path.exists(os.path.join(WEB_UPDATE_DIR, 'index.html')) and os.path.exists(WEB_UPDATE_VERSION_FILE):
            with open(WEB_UPDATE_VERSION_FILE, 'r', encoding='utf-8') as f:
                v = f.read().strip()
            if es_mas_nueva(v, WEB_VERSION):
                return v
    except Exception as e:
        print("Error leyendo versión de interfaz descargada:", e)
    return WEB_VERSION


def ruta_interfaz():
    """Dónde está el index.html que hay que abrir."""
    if es_mas_nueva(version_web_actual(), WEB_VERSION):
        return os.path.join(WEB_UPDATE_DIR, 'index.html')
    if getattr(sys, 'frozen', False):
        return os.path.join(sys._MEIPASS, 'app_calculadora', 'index.html')
    return os.path.join(os.path.dirname(__file__), '..', '.tmp', 'app_calculadora', 'index.html')


def _bajar(url, timeout=15):
    req = urllib.request.Request(url, headers={'User-Agent': 'Calculadora3D/' + APP_VERSION})
    with urllib.request.urlopen(req, timeout=timeout) as r:
        return r.read()


def _descargar_interfaz(info):
    """Baja los archivos de interfaz a una carpeta temporal y solo si TODOS
    llegan bien los mueve a su sitio. Así nunca queda una interfaz a medias."""
    archivos = info.get('web_files') or []
    if not archivos:
        return False
    temp = tempfile.mkdtemp(prefix='calc3d_web_')
    try:
        for nombre in archivos:
            datos = _bajar(WEB_BASE_URL + nombre, timeout=20)
            if not datos:
                raise ValueError('archivo vacío: ' + nombre)
            with open(os.path.join(temp, nombre), 'wb') as f:
                f.write(datos)

        if not os.path.exists(os.path.join(temp, 'index.html')):
            raise ValueError('el paquete no trae index.html')

        if not os.path.isdir(WEB_UPDATE_DIR):
            os.makedirs(WEB_UPDATE_DIR)
        for nombre in archivos:
            shutil.copy2(os.path.join(temp, nombre), os.path.join(WEB_UPDATE_DIR, nombre))
        with open(WEB_UPDATE_VERSION_FILE, 'w', encoding='utf-8') as f:
            f.write(str(info.get('web_version', '')))
        return True
    except Exception as e:
        print("Error descargando la interfaz:", e)
        return False
    finally:
        shutil.rmtree(temp, ignore_errors=True)


def _descargar_exe(url=None):
    """Baja el ejecutable indicado (por defecto el más reciente) al lado del
    actual, informando el avance."""
    global _estado_update
    url = url or EXE_URL
    destino = os.path.join(get_base_path(), 'Calculadora3D.nuevo.exe')
    parcial = destino + '.part'
    try:
        _estado_update = {'estado': 'descargando', 'porcentaje': 0, 'mensaje': ''}
        req = urllib.request.Request(url, headers={'User-Agent': 'Calculadora3D/' + APP_VERSION})
        with urllib.request.urlopen(req, timeout=30) as r:
            total = int(r.headers.get('Content-Length') or 0)
            bajado = 0
            with open(parcial, 'wb') as f:
                while True:
                    trozo = r.read(262144)
                    if not trozo:
                        break
                    f.write(trozo)
                    bajado += len(trozo)
                    if total:
                        _estado_update['porcentaje'] = int(bajado * 100 / total)

        if total and bajado < total:
            raise IOError('descarga incompleta')
        if bajado < 1000000:
            raise IOError('el archivo descargado es demasiado pequeño')

        if os.path.exists(destino):
            os.remove(destino)
        os.rename(parcial, destino)
        _estado_update = {'estado': 'listo', 'porcentaje': 100, 'mensaje': ''}
    except Exception as e:
        print("Error descargando el ejecutable:", e)
        try:
            if os.path.exists(parcial):
                os.remove(parcial)
        except Exception:
            pass
        _estado_update = {'estado': 'error', 'porcentaje': 0, 'mensaje': str(e)}


def _lanzar_reemplazo():
    """Windows no deja sobrescribir un .exe que está corriendo, así que un .bat
    espera a que la app cierre, cambia el archivo y la vuelve a abrir."""
    base = get_base_path()
    actual = sys.executable if getattr(sys, 'frozen', False) else os.path.join(base, 'Calculadora3D.exe')
    nuevo = os.path.join(base, 'Calculadora3D.nuevo.exe')
    if not os.path.exists(nuevo):
        return False

    bat = os.path.join(tempfile.gettempdir(), 'calc3d_actualizar.bat')
    contenido = (
        '@echo off\r\n'
        'set "ACTUAL={actual}"\r\n'
        'set "NUEVO={nuevo}"\r\n'
        'set "CARPETA={carpeta}"\r\n'
        ':esperar\r\n'
        'timeout /t 1 /nobreak >nul\r\n'
        '2>nul (>>"%ACTUAL%" call ) || goto esperar\r\n'
        'move /y "%NUEVO%" "%ACTUAL%" >nul\r\n'
        'timeout /t 1 /nobreak >nul\r\n'
        'cd /d "%CARPETA%"\r\n'
        'start "" "%ACTUAL%"\r\n'
        'del "%~f0"\r\n'
    ).format(actual=actual, nuevo=nuevo, carpeta=base)
    with open(bat, 'w', encoding='utf-8') as f:
        f.write(contenido)

    # PyInstaller le pasa a los procesos hijos variables como _MEIPASS2 y _PYI_*
    # que apuntan a la carpeta temporal de ESTA instancia. Si no se limpian, la
    # app relanzada intenta usar una carpeta que ya no existe y muere con "Error".
    entorno = {k: v for k, v in os.environ.items()
               if not (k.startswith('_MEI') or k.startswith('_PYI'))}

    # Solo CREATE_NO_WINDOW: con DETACHED_PROCESS el .bat se queda sin consola
    # y el comando "start" no llega a abrir la app de nuevo.
    subprocess.Popen(['cmd', '/c', bat],
                     cwd=base,
                     env=entorno,
                     creationflags=getattr(subprocess, 'CREATE_NO_WINDOW', 0))
    return True


def revisar_actualizaciones(window):
    """Corre en segundo plano al abrir la app. Si algo falla (sin internet,
    GitHub caído) simplemente no hace nada: la app funciona igual."""
    global _info_remota
    try:
        window.events.loaded.wait(30)
        time.sleep(1.5)
        info = json.loads(_bajar(VERSION_URL, timeout=10).decode('utf-8'))
    except Exception as e:
        print("Sin actualizaciones (no se pudo consultar):", e)
        return

    _info_remota = info
    prefs = cargar_prefs()

    # El usuario eligió quedarse en una versión concreta: no se toca nada
    if prefs.get('fijada'):
        print("Actualizaciones en pausa: el usuario fijó la versión", prefs['fijada'])
        return

    aviso = None

    # 1) Interfaz: se baja sola, se aplica al reiniciar
    try:
        if es_mas_nueva(info.get('web_version', '0'), version_web_actual()):
            if _descargar_interfaz(info):
                aviso = {
                    'tipo': 'web',
                    'version': info.get('web_version'),
                    'notas': info.get('notas', '')
                }
    except Exception as e:
        print("Error aplicando actualización de interfaz:", e)

    # 2) Ejecutable: solo se avisa, el usuario decide
    try:
        nueva_app = info.get('app_version', '0')
        if (getattr(sys, 'frozen', False)
                and es_mas_nueva(nueva_app, APP_VERSION)
                and str(nueva_app) not in prefs.get('omitidas', [])):
            aviso = {
                'tipo': 'exe',
                'version': nueva_app,
                'notas': info.get('notas', ''),
                'actual': APP_VERSION
            }
    except Exception as e:
        print("Error revisando versión del ejecutable:", e)

    if aviso:
        try:
            window.evaluate_js('window.avisoActualizacion && window.avisoActualizacion({})'.format(json.dumps(aviso)))
        except Exception as e:
            print("No se pudo mostrar el aviso de actualización:", e)
# ===== FIN ACTUALIZACIÓN AUTOMÁTICA =====

class Api:
    # ---- Actualizaciones (lo llama la interfaz) ----
    def info_version(self):
        prefs = cargar_prefs()
        return json.dumps({
            'app': APP_VERSION,
            'web': version_web_actual(),
            'repo': RELEASES_URL,
            'fijada': prefs.get('fijada'),
            'omitidas': prefs.get('omitidas', []),
            'empaquetada': bool(getattr(sys, 'frozen', False))
        })

    def listar_versiones(self):
        """Todas las versiones publicadas, para poder volver a una anterior."""
        try:
            datos = json.loads(_bajar(RELEASES_API + '?per_page=30', timeout=15).decode('utf-8'))
        except Exception as e:
            print("No se pudo consultar el listado de versiones:", e)
            return json.dumps({'error': 'No se pudo conectar con GitHub.'})

        prefs = cargar_prefs()
        lista = []
        for r in datos:
            if r.get('draft'):
                continue
            version = str(r.get('tag_name', '')).lstrip('vV')
            tiene_exe = any(a.get('name') == 'Calculadora3D.exe' for a in r.get('assets', []))
            if not version or not tiene_exe:
                continue
            lista.append({
                'version': version,
                'notas': (r.get('body') or '').split('\n\n')[0].strip(),
                'fecha': (r.get('published_at') or '')[:10],
                'instalada': version == APP_VERSION,
            })
        return json.dumps({
            'versiones': lista,
            'actual': APP_VERSION,
            'fijada': prefs.get('fijada'),
        })

    def omitir_version(self, version):
        prefs = cargar_prefs()
        version = str(version)
        if version not in prefs['omitidas']:
            prefs['omitidas'].append(version)
        guardar_prefs(prefs)
        return True

    def reactivar_actualizaciones(self):
        prefs = cargar_prefs()
        prefs['fijada'] = None
        prefs['omitidas'] = []
        guardar_prefs(prefs)
        return True

    def instalar_version(self, version):
        """Instala una versión concreta (normalmente una anterior). Deja la app
        fijada ahí para que el actualizador no la vuelva a mover sin permiso."""
        if _estado_update.get('estado') == 'descargando':
            return True
        version = str(version).lstrip('vV')

        def trabajo():
            _descargar_exe(url_exe_de_version(version))
            if _estado_update.get('estado') != 'listo':
                return
            prefs = cargar_prefs()
            prefs['fijada'] = version
            guardar_prefs(prefs)
            # La interfaz descargada puede ser más nueva que la de este .exe:
            # se borra para que quede exactamente la versión elegida.
            shutil.rmtree(WEB_UPDATE_DIR, ignore_errors=True)

        threading.Thread(target=trabajo, daemon=True).start()
        return True

    def iniciar_actualizacion_exe(self):
        if _estado_update.get('estado') == 'descargando':
            return True

        def trabajo():
            _descargar_exe()
            if _estado_update.get('estado') == 'listo':
                # Actualizar a propósito reactiva el seguimiento normal
                prefs = cargar_prefs()
                prefs['fijada'] = None
                guardar_prefs(prefs)

        threading.Thread(target=trabajo, daemon=True).start()
        return True

    def estado_actualizacion(self):
        return json.dumps(_estado_update)

    def aplicar_actualizacion_exe(self):
        if _estado_update.get('estado') != 'listo':
            return False
        if not _lanzar_reemplazo():
            return False
        # El .bat espera a que este proceso muera para cambiar el archivo
        threading.Thread(target=lambda: (time.sleep(0.4), webview.windows[0].destroy()), daemon=True).start()
        return True

    def abrir_descargas(self):
        try:
            webbrowser.open(RELEASES_URL)
            return True
        except Exception as e:
            print("No se pudo abrir la página de descargas:", e)
            return False

    def load_profiles(self):
        try:
            if os.path.exists(PROFILES_FILE):
                with open(PROFILES_FILE, 'r', encoding='utf-8') as f:
                    return f.read()
        except Exception as e:
            print("Error cargando perfiles:", e)
        return "{}"

    def save_profiles(self, data):
        try:
            parsed = json.loads(data)
            n_fil = len(parsed.get('filamentos', []))
            n_piezas = len(parsed.get('pieces', {}))
            log_path = os.path.join(get_base_path(), 'save_log.txt')
            with open(log_path, 'a', encoding='utf-8') as log:
                log.write(f"[{datetime.datetime.now()}] filamentos={n_fil}, piezas={n_piezas}\n")
                for f in parsed.get('filamentos', []):
                    log.write(f"  -> {f.get('marca')} {f.get('tipo')} {f.get('color')}\n")
            hacer_copia()
            with open(PROFILES_FILE, 'w', encoding='utf-8') as f:
                f.write(data)
            return True
        except Exception as e:
            print("Error guardando perfiles:", e)
            return False

    # ---- Copias de seguridad (lo llama la ventana de Ajustes) ----
    def listar_copias(self):
        salida = []
        try:
            for archivo in _copias_ordenadas():
                ruta = os.path.join(COPIAS_DIR, archivo)
                marca = archivo.replace('perfiles_', '').replace('.json', '')
                fecha, hora = (marca.split('_') + [''])[:2]
                kb = os.path.getsize(ruta) / 1024.0
                salida.append({
                    'archivo': archivo,
                    'fecha': '{} a las {}'.format(fecha, hora.replace('-', ':')),
                    'tamano': '{:.0f} KB'.format(kb) if kb >= 1 else '<1 KB',
                })
        except Exception as e:
            print("Error listando copias:", e)
        return json.dumps(salida)

    def restaurar_copia(self, archivo):
        try:
            # Solo un nombre de archivo, nunca una ruta: evita salirse de la carpeta
            archivo = os.path.basename(str(archivo))
            origen = os.path.join(COPIAS_DIR, archivo)
            if not os.path.exists(origen):
                return False
            with open(origen, 'r', encoding='utf-8') as f:
                contenido = f.read()
            json.loads(contenido)          # si la copia está corrupta, no se restaura
            hacer_copia()                  # lo de ahora también se guarda por si acaso
            with open(PROFILES_FILE, 'w', encoding='utf-8') as f:
                f.write(contenido)
            return True
        except Exception as e:
            print("Error restaurando la copia:", e)
            return False

    def export_profiles_file(self, data):
        try:
            file_types = ('Archivos JSON (*.json)',)
            result = webview.windows[0].create_file_dialog(webview.SAVE_DIALOG, file_types=file_types, save_filename='mis_perfiles_3d.json')
            if not result:
                return False
            path = result[0]
            if not path.endswith('.json'):
                path += '.json'
            with open(path, 'w', encoding='utf-8') as f:
                f.write(data)
            return True
        except Exception as e:
            print("Error exportando perfiles:", e)
            return False

    def import_profiles_file(self):
        try:
            file_types = ('Archivos JSON (*.json)',)
            result = webview.windows[0].create_file_dialog(webview.OPEN_DIALOG, file_types=file_types)
            if not result:
                return "{}"
            path = result[0]
            with open(path, 'r', encoding='utf-8') as f:
                return f.read()
        except Exception as e:
            print("Error importando perfiles:", e)
            return "{}"

    def _dibujar_empresa(self, pdf, empresa_json):
        """Membrete con los datos de la empresa. Devuelve la nota al pie (si hay)
        para imprimirla al final. Todos los campos son opcionales."""
        if not empresa_json:
            return ''
        try:
            emp = json.loads(empresa_json) if isinstance(empresa_json, str) else empresa_json
        except Exception:
            return ''
        if not isinstance(emp, dict) or not emp:
            return ''

        logo_temp = None
        try:
            logo = emp.get('logo') or ''
            if logo.startswith('data:image'):
                import base64
                cabecera, datos64 = logo.split(',', 1)
                ext = '.png' if 'png' in cabecera else '.jpg'
                fd, logo_temp = tempfile.mkstemp(suffix=ext, prefix='calc3d_logo_')
                with os.fdopen(fd, 'wb') as f:
                    f.write(base64.b64decode(datos64))
                pdf.image(logo_temp, x=10, y=8, h=18)
        except Exception as e:
            print("No se pudo poner el logo en el PDF:", e)

        try:
            pdf.set_xy(35, 10)
            pdf.set_font("Arial", 'B', 13)
            pdf.cell(0, 6, txt=_texto_pdf(emp.get('nombre', '')), ln=1)
            pdf.set_font("Arial", size=9)
            renglones = []
            if emp.get('nit'):
                renglones.append('NIT: ' + emp['nit'])
            contacto = ' | '.join(x for x in [emp.get('telefono'), emp.get('email'), emp.get('web')] if x)
            if contacto:
                renglones.append(contacto)
            if emp.get('direccion'):
                renglones.append(emp['direccion'])
            for r in renglones:
                pdf.set_x(35)
                pdf.cell(0, 4.5, txt=_texto_pdf(r), ln=1)

            y = max(pdf.get_y(), 28)
            pdf.line(10, y + 1, 200, y + 1)
            pdf.set_y(y + 5)
        except Exception as e:
            print("No se pudo poner el membrete:", e)
        finally:
            if logo_temp:
                try:
                    os.remove(logo_temp)
                except Exception:
                    pass

        return emp.get('nota', '') or ''

    def _nota_pie(self, pdf, nota):
        if not nota:
            return
        try:
            pdf.ln(8)
            pdf.set_font("Arial", 'I', 9)
            pdf.set_text_color(110, 110, 110)
            pdf.multi_cell(0, 4.5, txt=_texto_pdf(nota))
            pdf.set_text_color(0, 0, 0)
        except Exception as e:
            print("No se pudo poner la nota al pie:", e)

    def export_pdf(self, total, filament, tiempo_prod, nombre='', empresa_json=''):
        try:
            file_types = ('Archivos PDF (*.pdf)',)
            fname = f'Cotizacion ({nombre}).pdf' if nombre else 'Cotizacion_Impresion.pdf'
            result = webview.windows[0].create_file_dialog(webview.SAVE_DIALOG, file_types=file_types, save_filename=fname)
            if not result:
                return False
                
            path = result[0]
            if not path.endswith('.pdf'):
                path += '.pdf'
                
            pdf = FPDF()
            pdf.add_page()
            nota_pie = self._dibujar_empresa(pdf, empresa_json)
            pdf.set_font("Arial", 'B', 16)
            pdf.cell(200, 10, txt=_texto_pdf("Cotización Impresión"), ln=1, align='C')
            pdf.ln(10)
            pdf.set_font("Arial", size=12)
            date_str = datetime.datetime.now().strftime("%d/%m/%Y %H:%M")
            pdf.cell(200, 10, txt=f"Fecha: {date_str}", ln=1, align='L')
            if nombre:
                pdf.cell(200, 10, txt=_texto_pdf(f"Pieza: {nombre}"), ln=1, align='L')
            pdf.ln(5)
            pdf.set_font("Arial", 'B', 12)
            pdf.cell(100, 10, txt="Gramos de Filamento Estimados: ", ln=0)
            pdf.set_font("Arial", size=12)
            pdf.cell(100, 10, txt=f"{filament} g", ln=1)
            
            if tiempo_prod:
                pdf.set_font("Arial", 'B', 12)
                pdf.cell(100, 10, txt="Tiempo Est. de Producción (Aprox): ", ln=0)
                pdf.set_font("Arial", size=12)
                pdf.cell(100, 10, txt=f"{tiempo_prod}", ln=1)

            pdf.ln(5)
            pdf.set_font("Arial", 'B', 14)
            pdf.cell(100, 10, txt="Costo Total a Cobrar: ", ln=0)
            pdf.set_font("Arial", size=14)
            pdf.set_text_color(20, 160, 40)
            pdf.cell(100, 10, txt=f"{total}", ln=1)
            pdf.set_text_color(0, 0, 0)

            self._nota_pie(pdf, nota_pie)
            pdf.output(path)
            return True
        except Exception as e:
            print("Error exportando PDF:", e)
            return False

    def export_excel(self, data_json, nombre=''):
        try:
            data = json.loads(data_json)
            file_types = ('Archivos Excel (*.xlsx)',)
            fname = f'Cotizacion ({nombre}).xlsx' if nombre else 'Desglose_Costos.xlsx'
            result = webview.windows[0].create_file_dialog(webview.SAVE_DIALOG, file_types=file_types, save_filename=fname)
            if not result:
                return False
            path = result[0]
            if not path.endswith('.xlsx'):
                path += '.xlsx'
            wb = Workbook()
            ws = wb.active
            ws.title = "Costos"
            ws.append(["Concepto", "Valor"])
            for key, value in data.items():
                ws.append([key, value])
            wb.save(path)
            return True
        except Exception as e:
            print("Error exportando Excel:", e)
            return False

    def export_bom(self, report, nombre):
        try:
            file_types = ('Archivos de texto (*.txt)',)
            fname = f'BOM_{nombre.replace(" ", "_")}.txt' if nombre else 'BOM_Lista_Materiales.txt'
            result = webview.windows[0].create_file_dialog(webview.SAVE_DIALOG, file_types=file_types, save_filename=fname)
            if not result:
                return False
                
            path = result[0]
            if not path.endswith('.txt'):
                path += '.txt'
                
            with open(path, 'w', encoding='utf-8') as f:
                f.write(report)
            return True
        except Exception as e:
            print("Error exportando BOM:", e)
            return False

    def export_project_pdf(self, total_proyecto, gramos_totales, camas, unids_pedido, unids_cama, precio_por_cama, costo_materiales, insumos_extra, mano_obra, tiempo_proy, costo_unidad_pdf, nombre='', costo_prod_pieza=None, costo_prod_total=None, empresa_json=''):
        try:
            file_types = ('Archivos PDF (*.pdf)',)
            fname = f'Cotizacion ({nombre}).pdf' if nombre else 'Cotizacion_Proyecto.pdf'
            result = webview.windows[0].create_file_dialog(webview.SAVE_DIALOG, file_types=file_types, save_filename=fname)
            if not result:
                return False
                
            path = result[0]
            if not path.endswith('.pdf'):
                path += '.pdf'
                
            pdf = FPDF()
            pdf.add_page()
            nota_pie = self._dibujar_empresa(pdf, empresa_json)
            pdf.set_font("Arial", 'B', 16)
            pdf.cell(200, 10, txt=_texto_pdf("Cotización de Proyecto 3D"), ln=1, align='C')
            pdf.ln(10)
            
            pdf.set_font("Arial", size=12)
            date_str = datetime.datetime.now().strftime("%d/%m/%Y %H:%M")
            pdf.cell(200, 10, txt=f"Fecha: {date_str}", ln=1, align='L')
            pdf.ln(5)
            
            pdf.set_font("Arial", 'B', 12)
            pdf.cell(100, 10, txt="Detalles del Pedido:", ln=1)
            pdf.set_font("Arial", size=12)
            pdf.cell(100, 10, txt=f"Unidades Totales: {unids_pedido}", ln=1)
            pdf.cell(100, 10, txt=f"Unidades por Cama: {unids_cama}", ln=1)
            pdf.cell(100, 10, txt=f"Total de Camas (Impresiones): {camas}", ln=1)
            pdf.cell(100, 10, txt=f"Gramos de Filamento Totales Estimados: {gramos_totales} g", ln=1)
            if tiempo_proy:
                pdf.cell(100, 10, txt=f"Tiempo Est. de Producción (Aprox): {tiempo_proy}", ln=1)
            
            pdf.ln(5)
            
            pdf.set_font("Arial", 'B', 12)
            pdf.cell(100, 10, txt="Desglose Financiero del Proyecto:", ln=1)
            pdf.set_font("Arial", size=12)
            if insumos_extra and insumos_extra not in ["0", "$ 0", "COP 0", "USD 0"]:
                pdf.cell(100, 10, txt=f"Insumos Extra: {insumos_extra}", ln=1)
            pdf.cell(100, 10, txt=f"Costo por Unidad: {costo_unidad_pdf}", ln=1)
            if costo_prod_pieza:
                pdf.cell(100, 10, txt=f"Costo Produccion (1 pieza): {costo_prod_pieza}", ln=1)
            if costo_prod_total:
                pdf.cell(100, 10, txt=f"Costo Produccion (proyecto completo): {costo_prod_total}", ln=1)
            pdf.ln(5)
            
            pdf.set_font("Arial", 'B', 14)
            pdf.cell(100, 10, txt="Costo Total a Cobrar por Proyecto: ", ln=0)
            pdf.set_font("Arial", size=14)
            pdf.set_text_color(20, 160, 40)
            pdf.cell(100, 10, txt=f"{total_proyecto}", ln=1)
            pdf.set_text_color(0, 0, 0)

            self._nota_pie(pdf, nota_pie)
            pdf.output(path)
            return True
        except Exception as e:
            print("Error exportando PDF de Proyecto:", e)
            return False

    def export_project_excel(self, data_json, nombre=''):
        try:
            data = json.loads(data_json)
            file_types = ('Archivos Excel (*.xlsx)',)
            fname = f'Cotizacion ({nombre}).xlsx' if nombre else 'Cotizacion_Proyecto.xlsx'
            result = webview.windows[0].create_file_dialog(webview.SAVE_DIALOG, file_types=file_types, save_filename=fname)
            if not result:
                return False
            path = result[0]
            if not path.endswith('.xlsx'):
                path += '.xlsx'
            wb = Workbook()
            ws = wb.active
            ws.title = "Proyecto"
            ws.append(["Concepto", "Valor"])
            for key, value in data.items():
                ws.append([key, value])
            wb.save(path)
            return True
        except Exception as e:
            print("Error exportando Excel de Proyecto:", e)
            return False

    def export_megaproject_pdf(self, payload, empresa_json=''):
        try:
            data = json.loads(payload)
            file_types = ('Archivos PDF (*.pdf)',)
            mp_nombre = data.get('nombre', '')
            fname = f'Cotizacion ({mp_nombre}).pdf' if mp_nombre and mp_nombre != 'Mega Proyecto' else 'Cotizacion_MegaProyecto.pdf'
            result = webview.windows[0].create_file_dialog(webview.SAVE_DIALOG, file_types=file_types, save_filename=fname)
            if not result:
                return False
                
            path = result[0]
            if not path.endswith('.pdf'):
                path += '.pdf'
                
            pdf = FPDF()
            pdf.add_page()
            nota_pie = self._dibujar_empresa(pdf, empresa_json)
            pdf.set_font("Arial", 'B', 16)
            nombre_proyecto = data.get('nombre', 'Mega Proyecto')
            pdf.cell(200, 10, txt=_texto_pdf(f"Cotización - {nombre_proyecto}"), ln=1, align='C')
            pdf.ln(10)
            
            pdf.set_font("Arial", size=12)
            date_str = datetime.datetime.now().strftime("%d/%m/%Y %H:%M")
            pdf.cell(200, 10, txt=f"Fecha: {date_str}", ln=1, align='L')
            pdf.ln(5)
            
            pdf.set_font("Arial", 'B', 14)
            pdf.cell(200, 10, txt="Desglose por Proyecto:", ln=1)
            pdf.ln(2)
            
            # List items
            for item in data.get('items', []):
                pdf.set_font("Arial", 'B', 12)
                pdf.cell(0, 8, txt=f"- {item.get('nombre', '')} ({item.get('unidades', '')} u)", ln=1)
                pdf.set_font("Arial", size=10)
                pdf.cell(0, 6, txt=f"  Camas: {item.get('camas', '')} | Insumos: {item.get('insumos', '')} | Tiempo Prod: {item.get('tiempo', '')}", ln=1)
                costo_prod_str = f" | Costo Prod. c/u: {item.get('costo_prod_pieza', '')}" if item.get('costo_prod_pieza') else ""
                pdf.cell(0, 6, txt=f"  Precio por unidad: {item.get('costo_unidad', '')} | Precio a cobrar: {item.get('total', '')}{costo_prod_str}", ln=1)
                pdf.ln(3)
                
            pdf.ln(5)
            
            # Totals
            pdf.set_font("Arial", 'B', 14)
            pdf.cell(200, 10, txt="TOTALES DEL MEGA PROYECTO:", ln=1)
            
            totales = data.get('totales', {})
            pdf.set_font("Arial", size=12)
            pdf.cell(100, 8, txt=f"Total de Camas: {totales.get('camas', '')}", ln=1)
            pdf.cell(100, 8, txt=f"Horas: {totales.get('horas', '')} h", ln=1)
            if totales.get('tiempo'):
                pdf.cell(100, 8, txt=f"Tiempo Estimado Produccion: {totales.get('tiempo', '')}", ln=1)
            
            pdf.ln(3)
            if totales.get('insumos') and totales.get('insumos') not in ["0", "$ 0", "COP 0", "USD 0"]:
                pdf.cell(100, 8, txt=f"Insumos: {totales.get('insumos', '')}", ln=1)
            if totales.get('costo_prod_pieza'):
                pdf.cell(100, 8, txt=f"Costo Produccion promedio (c/u): {totales.get('costo_prod_pieza', '')}", ln=1)
            if totales.get('costo_prod_total'):
                pdf.cell(100, 8, txt=f"Costo Produccion total Mega Proyecto: {totales.get('costo_prod_total', '')}", ln=1)
            
            pdf.ln(5)
            pdf.set_font("Arial", 'B', 16)
            pdf.set_text_color(20, 160, 40)
            pdf.cell(100, 10, txt=_texto_pdf(f"TOTAL A COBRAR MEGA PROYECTO: {totales.get('total', '')}"), ln=1)
            pdf.set_text_color(0, 0, 0)

            self._nota_pie(pdf, nota_pie)
            pdf.output(path)
            return True
        except Exception as e:
            print("Error exportando PDF Mega Proyecto:", e)
            return False

    def export_megaproject_excel(self, payload):
        try:
            data = json.loads(payload)
            file_types = ('Archivos Excel (*.xlsx)',)
            mp_nombre = data.get('nombre', '')
            fname = f'Cotizacion ({mp_nombre}).xlsx' if mp_nombre and mp_nombre != 'Mega Proyecto' else 'Cotizacion_MegaProyecto.xlsx'
            result = webview.windows[0].create_file_dialog(webview.SAVE_DIALOG, file_types=file_types, save_filename=fname)
            if not result:
                return False
            path = result[0]
            if not path.endswith('.xlsx'):
                path += '.xlsx'
            wb = Workbook()
            ws = wb.active
            ws.title = "Mega Proyecto"
            
            ws.append(["--- DETALLE POR PROYECTO ---"])
            if data.get('items'):
                headers = list(data['items'][0].keys())
                ws.append(headers)
                for item in data['items']:
                    ws.append([item.get(h, "") for h in headers])
                    
            ws.append([])
            ws.append(["--- TOTALES AGREGADOS ---"])
            totales = data.get('totales', {})
            for key, value in totales.items():
                ws.append([key, value])
                
            wb.save(path)
            return True
        except Exception as e:
            print("Error exportando Excel de Mega Proyecto:", e)
            return False

if __name__ == '__main__':
    # Si ya se descargó una interfaz más nueva, se abre esa en vez de la incluida
    html_path = ruta_interfaz()

    api = Api()
    win_state = load_window_state()

    window = webview.create_window(
        title='Calculadora de Impresión 3D - by samderbrow',
        url=f'file:///{html_path.replace(chr(92), "/")}',
        js_api=api,
        width=win_state['width'],
        height=win_state['height'],
        x=win_state['x'],
        y=win_state['y'],
        maximized=win_state['maximized'],
        min_size=(MIN_W, MIN_H)
    )

    # Estado actual: solo se registra el tamaño/posición "normal" (no maximizado),
    # que es el que se debe restaurar al des-maximizar en la próxima sesión.
    current = dict(win_state)

    def is_maximized_now():
        # Los eventos de maximizar/redimensionar llegan en hilos distintos, así que
        # se consulta el estado real de la ventana antes de guardar medidas.
        try:
            form = window.native
            if form is not None:
                return str(form.WindowState) != 'Normal'
        except Exception:
            pass
        return current['maximized']

    def on_resized(width, height):
        if is_maximized_now():
            return
        if width >= MIN_W and height >= MIN_H:
            current['width'], current['height'] = int(width), int(height)

    def on_moved(x, y):
        if is_maximized_now():
            return
        # Al minimizar, Windows reporta posiciones tipo -32000
        if x > -10000 and y > -10000:
            current['x'], current['y'] = int(x), int(y)

    def on_maximized():
        current['maximized'] = True

    def on_restored():
        current['maximized'] = False

    def native_snapshot():
        """Geometría exacta leída de la ventana nativa (Windows/WinForms).

        RestoreBounds da el tamaño "normal" incluso estando maximizada, que es
        justo lo que hay que recordar. Si falla, se usa lo captado por eventos.
        """
        try:
            form = window.native
            if form is None:
                return None
            scale = getattr(form, '_scale', 1) or 1
            state = str(form.WindowState)
            # Con la ventana en estado normal manda su posición actual; maximizada o
            # minimizada hay que mirar RestoreBounds (el tamaño al des-maximizar).
            b = form.Bounds if state == 'Normal' else form.RestoreBounds
            w, h = int(b.Width / scale), int(b.Height / scale)
            if w < MIN_W or h < MIN_H:
                return None
            return {
                'width': w,
                'height': h,
                'x': int(b.X / scale),
                'y': int(b.Y / scale),
                # Si está minimizada no se sabe el estado previo: se usa el rastreado
                'maximized': True if state == 'Maximized' else (
                    current['maximized'] if state == 'Minimized' else False
                )
            }
        except Exception as e:
            print("No se pudo leer la geometría nativa:", e)
            return None

    def on_closing():
        save_window_state(native_snapshot() or current)

    window.events.resized += on_resized
    window.events.moved += on_moved
    window.events.maximized += on_maximized
    window.events.restored += on_restored
    window.events.closing += on_closing

    threading.Thread(target=revisar_actualizaciones, args=(window,), daemon=True).start()

    webview.start(private_mode=False)
